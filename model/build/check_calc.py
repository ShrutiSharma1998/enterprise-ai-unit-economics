# SPDX-License-Identifier: Apache-2.0
"""Stage check: build inputs + calc sheets, recalc in Excel, compare every total with the reference model."""
import sys
import openpyxl
import build_inputs as bi
import build_calc as bc
import reference_model as ref
from xl_run import save_and_recalc, errors

OUT = "_stage2.xlsx"


def build():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bi.build_org(wb)
    bi.build_params(wb)
    bi.build_pricebook(wb)
    bi.build_teams(wb)
    for k in range(3):
        bc.build_calc(wb, k)
    return wb


def compare(v, tol=1e-6):
    bad, n = [], 0
    for k in range(3):
        ws = v[bc.NAMES[k]]
        r = ref.run(k)
        for t in range(bi.NT):
            for name in bc.COST_KEYS:
                got = ws[f"C{bc.ROWS['team'](t, name)}"].value or 0
                exp = r["teams"][t]["direct"][name]
                n += 1
                if abs(got - exp) > tol * max(1, abs(exp)):
                    bad.append((k, t, name, got, exp))
            for name, exp in [("benefit", r["teams"][t]["benefit"]), ("out", r["teams"][t]["outcomes"]), ("act", r["teams"][t]["active_months"])]:
                got = ws[f"C{bc.ROWS['team'](t, name)}"].value or 0
                n += 1
                if abs(got - exp) > tol * max(1, abs(exp)):
                    bad.append((k, t, name, got, exp))
            got = ws[f"C{bc.ROWS['team'](t, 'pb')}"].value
            exp = r["teams"][t]["payback"] if r["teams"][t]["payback"] is not None else "None"
            n += 1
            if got != exp:
                bad.append((k, t, "payback", got, exp))
            got = ws[f"{bc.ALLOC_COLS[t]}{bc.ROWS['alloc']('alloc')}"].value
            n += 1
            if abs(got - r["alloc"][t]) > tol * max(1, abs(got)):
                bad.append((k, t, "alloc", got, r["alloc"][t]))
            got = ws[f"C{bc.ROWS['team'](t, 'net')}"].value
            exp = r["teams"][t]["net"]
            n += 1
            if abs(got - exp) > 1e-6 * max(1, abs(exp)):
                bad.append((k, t, "net(after alloc)", got, exp))
        for name in bc.HUB_COST_KEYS:
            got = ws[f"C{bc.ROWS['hub'](name)}"].value or 0
            exp = r["hub_ot"].get(name, None)
            if exp is None:
                exp = sum(r["hub_rows"][name])
            n += 1
            if abs(got - exp) > tol * max(1, abs(exp)):
                bad.append((k, "hub", name, got, exp))
        for label, got, exp in [
            ("ent benefit", ws[f"C{bc.ROWS['ent']('benefit')}"].value, r["comp"]["benefit"]),
            ("ent net", ws[f"C{bc.ROWS['ent']('net')}"].value, r["net"]),
            ("ent payback", ws[f"C{bc.ROWS['ent']('pb')}"].value, r["payback"] if r["payback"] is not None else "None"),
            ("unalloc", ws[f"{bc.ALLOC_UNALLOC}{bc.ROWS['alloc']('alloc')}"].value, r["unalloc"]),
        ]:
            n += 1
            if isinstance(exp, str) or isinstance(got, str):
                if got != exp:
                    bad.append((k, "ent", label, got, exp))
            elif abs(got - exp) > tol * max(1, abs(exp)):
                bad.append((k, "ent", label, got, exp))
        for t in range(bi.NT):
            got = ws[f"C{bc.ROWS['team'](t, 'hdr')}"].value
            n += 1
            if abs(got - r["cpo"][t]) > 1e-9:
                bad.append((k, t, "cost per outcome", got, r["cpo"][t]))
    return n, bad


if __name__ == "__main__":
    v = save_and_recalc(build(), OUT)
    print("formula errors:", errors(v)[:5])
    n, bad = compare(v)
    print(f"compared {n} values; mismatches: {len(bad)}")
    for b in bad[:25]:
        print("  ", b)
    sys.exit(1 if bad else 0)
