# SPDX-License-Identifier: Apache-2.0
"""Phase 5 validation. Reads the built workbook, and re-runs Excel for the sensitivity linearity check.

  python validate.py envelopes            envelope checks against published and derived ranges (no Excel run)
  python validate.py bands                scenario band decomposition (no Excel run)
  python validate.py sensitivity <tmp>    change each driver for real, recalculate, compare with the sheet's linear prediction
"""
import os
import sys
import openpyxl
import build as B
import build_inputs as bi
import build_calc as bc
import build_out_a as oa
import build_out_b as ob
from spec_data import TEAMS
from xl_run import save_and_recalc, cell

WB = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "ai-unit-economics-v0.1.xlsx")


def load():
    B.build()                                   # populates the address registries; the file itself is not rebuilt
    return openpyxl.load_workbook(WB, data_only=True)


def envelopes():
    v = load()
    res, RES = v["Results"], oa.RES
    hz = v["Org"][cell(bi.ADDR["org"]["horizon"])].value
    yrs = hz / 12
    print(f"Base scenario, horizon {hz} months. Per licensed user per year, USD.")
    print("team".ljust(30), "seats", "tokens+vendor+tools", "human review", "all other run+setup", "TOTAL incl hub", "| IT view (excl. review)")
    for t in range(bi.NT):
        c = oa.TCOL[t]
        L = TEAMS[t]["licensed"]
        g = lambda key: res[f"{c}{RES['cost_row'][key]}"].value or 0
        seat = g("seat")
        infer = g("tokens") + g("vendor") + g("tools")
        rev = g("review")
        total = res[f"{c}{RES['m']['total']}"].value
        other = total - seat - infer - rev
        d = L * yrs
        print(TEAMS[t]["name"][:29].ljust(30), f"{seat / d:7,.0f} {infer / d:10,.0f} {rev / d:14,.0f} {other / d:16,.0f} {total / d:14,.0f} | {(total - rev) / d:8,.0f}")
    print("\nSeat-only arithmetic: blended seat price x 12")
    for t in range(bi.NT):
        c = oa.TCOL[t]
    v_seat = v["Teams"]
    for t in range(bi.NT):
        r = bi.ADDR["derived_row"]["v_seat"]
        print("  ", TEAMS[t]["name"][:29].ljust(30), f"${v_seat[f'{bi.ADDR['team_cols'][t]}{r}'].value * 12:,.0f} per licensed seat per year (billed only when seats apply)")
    print("\nOne-time cost of team 2 (API-based coding-assistant-like) for the Gartner initial-cost envelope ($100k to $200k):")
    ws = v["Calc_Base"]
    ot = sum(ws[f"D{bc.ROWS['team'](1, k)}"].value or 0 for k in ("ot_data_effort", "ot_data_load", "ot_l2", "ot_train"))
    hub_ot = ws[f"D{bc.ROWS['hub']('raw')}"].value
    share = ws[f"C{bc.ROWS['alloc']('f')}"].value
    print(f"   team one-time ${ot:,.0f}; hub one-time (all) ${hub_ot:,.0f}; team 2 share of hub {share:.1%} = ${hub_ot * share:,.0f}; total ${ot + hub_ot * share:,.0f}")
    print("\nCost per outcome and per active user (after the horizon-total fix):")
    for t in range(bi.NT):
        c = oa.TCOL[t]
        print("  ", TEAMS[t]["name"][:29].ljust(30), "outcomes", f"{res[f'{c}{RES['m']['outcomes']}'].value:,.0f}",
              "cost/outcome", f"{res[f'{c}{RES['m']['cpo']}'].value:.2f}", "benefit/outcome", f"{res[f'{c}{RES['m']['bpo']}'].value:.2f}",
              "cost/active-user-month", res[f"{c}{RES['m']['cpa']}"].value if isinstance(res[f"{c}{RES['m']['cpa']}"].value, str) else f"{res[f'{c}{RES['m']['cpa']}'].value:.2f}")
    I = oa.ENTC
    print("   enterprise".ljust(33), "outcomes", f"{res[f'{I}{RES['m']['outcomes']}'].value:,.0f}", "cost/outcome", f"{res[f'{I}{RES['m']['cpo']}'].value:.2f}",
          "benefit/outcome", f"{res[f'{I}{RES['m']['bpo']}'].value:.2f}", "cost/active-user-month (per-user teams)", f"{res[f'{I}{RES['m']['cpa']}'].value:.2f}")
    print("   totals: cost", f"{res[f'{oa.ENTC}{RES['total_row']}'].value:,.0f}", "benefit", f"{res[f'{oa.ENTC}{RES['m']['benefit']}'].value:,.0f}", "net", f"{res[f'{oa.ENTC}{RES['m']['net']}'].value:,.0f}",
          "payback", res[f"{oa.ENTC}{RES['m']['payback']}"].value)


def bands():
    v = load()
    s = v["Scenarios"]
    print("Scenario totals (enterprise, over the horizon):")
    for lab, row in [("Total cost", ob.SCN["total"]), ("Benefit", ob.SCN["benefit"]), ("Net value", ob.SCN["net"])]:
        lo, ba, hi = (s[f"{c}{row}"].value for c in "BCD")
        print(f"  {lab:11s} low {lo:>14,.0f}  base {ba:>14,.0f}  high {hi:>14,.0f}   low/base {lo / ba:5.2f}x  high/base {hi / ba:5.2f}x" if ba else lab)
    for n in range(1, 6):
        row = ob.SCN["layer_row"][n]
        lo, ba, hi = (s[f"{c}{row}"].value for c in "BCD")
        print(f"  L{n}          low {lo:>14,.0f}  base {ba:>14,.0f}  high {hi:>14,.0f}   high/base {hi / ba:5.2f}x")
    tot = [s[f"{c}{ob.SCN['total']}"].value for c in "BCD"]
    l5 = [s[f"{c}{ob.SCN['layer_row'][5]}"].value for c in "BCD"]
    print("  share of the high-cost increase that is L5 (people and change):", f"{(l5[2] - l5[1]) / (tot[2] - tot[1]):.0%}")
    print("  published forecast-miss context (FCS-01, grade C): 24% of firms miss by more than 50%; 80% miss infrastructure forecasts by more than 25%")


def scale_teams(wb, key, factor, only=None):
    for t in range(bi.NT):
        if only is not None and t not in only:
            continue
        addr = f"{bi.ADDR['team_cols'][t]}{bi.ADDR['team_row'][key]}"
        wb["Teams"][addr].value = TEAMS[t][key] * factor


def sensitivity(tmp):
    base = load()
    sen = base["Sensitivity"]
    r0 = ob.SENS["r0"]
    N0 = sen["B6"].value
    print(f"Base net value N0 = {N0:,.0f}")
    drivers = [
        (0, "Realisation share", 1.5, lambda wb: wb["Params"].__setitem__(f"E{bi.ADDR['param_row']['realise']}", 0.5 * 1.5)),
        (1, "Minutes saved per outcome", 1.5, lambda wb: scale_teams(wb, "min_saved", 1.5)),
        (2, "Acceptance rate", 0.8, lambda wb: scale_teams(wb, "accept", 0.8)),
        (3, "Adoption and volume", 0.7, lambda wb: (scale_teams(wb, "plateau", 0.7), scale_teams(wb, "direct", 0.7))),
        (4, "Tokens per outcome", 1.5, lambda wb: (scale_teams(wb, "tok_in", 1.5), scale_teams(wb, "tok_out", 1.5))),
        (5, "Human review minutes", 2.0, lambda wb: wb["Params"].__setitem__(f"E{bi.ADDR['param_row']['rev_min']}", 3 * 2.0)),
        (6, "Effort (person-day rate x2)", 2.0, lambda wb: wb["Org"].__setitem__(cell(bi.ADDR["org"]["pd_rate"]), 1600)),
    ]
    print("driver".ljust(30), "k".rjust(5), "predicted net".rjust(15), "recalculated net".rjust(17), "difference".rjust(13), "error vs swing")
    for i, name, k, mutate in drivers:
        E = sen[f"B{r0 + i}"].value
        pred = N0 + (k - 1) * E
        wb = B.build()
        mutate(wb)
        v = save_and_recalc(wb, os.path.join(tmp, f"sens_{i}.xlsx"))
        got = v["Results"][f"{oa.ENTC}{oa.RES['m']['net']}"].value
        swing = abs((k - 1) * E) or 1
        print(name.ljust(30), f"{k:5.2f}", f"{pred:15,.0f}", f"{got:17,.0f}", f"{got - pred:13,.0f}", f"{abs(got - pred) / swing:8.1%}")


if __name__ == "__main__":
    mode = sys.argv[1]
    if mode == "envelopes":
        envelopes()
    elif mode == "bands":
        bands()
    else:
        os.makedirs(sys.argv[2], exist_ok=True)
        sensitivity(sys.argv[2])
