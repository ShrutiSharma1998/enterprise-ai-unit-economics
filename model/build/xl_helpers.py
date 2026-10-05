# SPDX-License-Identifier: Apache-2.0
"""Small styling and addressing helpers shared by the builder modules."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

FONT = "Arial"
BLUE, BLACK, GREEN, GREY, WHITE, RED = "0000FF", "000000", "008000", "666666", "FFFFFF", "C00000"
F_IN = PatternFill("solid", fgColor="FFF9D6")      # input cells (pale yellow)
F_HEAD = PatternFill("solid", fgColor="1F3864")    # header (dark blue)
F_SUB = PatternFill("solid", fgColor="D9E2F3")     # sub-header
F_NOTE = PatternFill("solid", fgColor="F2F2F2")
F_WARN = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

USD = '$#,##0;($#,##0);-'
USD2 = '$#,##0.00;($#,##0.00);-'
USD4 = '$#,##0.0000;($#,##0.0000);-'
NUM = '#,##0;(#,##0);-'
NUM2 = '#,##0.00;(#,##0.00);-'
PCT = '0.0%;(0.0%);-'
MULT = '0.00"x"'

MONTHS = 60
C0 = 4            # column D = month 0
C1 = 5            # column E = month 1
CN = C1 + MONTHS - 1   # last month column (BL)


def mc(m):
    """Column letter for month m (0..60)."""
    return L(C0 if m == 0 else C1 + m - 1)


def font(color=BLACK, bold=False, italic=False, size=10):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


def put(ws, ref, value, kind="text", fmt=None, bold=False, fill=None, wrap=False, italic=False, align=None):
    """kind: text | input | calc | link | head | sub | note"""
    c = ws[ref]
    c.value = value
    color = {"input": BLUE, "link": GREEN, "head": WHITE, "note": GREY}.get(kind, BLACK)
    c.font = font(color, bold=bold or kind == "head", italic=italic or kind == "note")
    if kind == "input":
        c.fill = F_IN
        c.border = BOX
    if kind == "head":
        c.fill = F_HEAD
    if kind == "sub":
        c.fill = F_SUB
        c.font = font(BLACK, bold=True)
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if wrap or align:
        c.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
    return c


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def title(ws, text, sub=None):
    put(ws, "A1", text, bold=True)
    ws["A1"].font = font(BLACK, bold=True, size=14)
    if sub:
        put(ws, "A2", sub, "note")
